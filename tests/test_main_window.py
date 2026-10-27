import copy
import csv
import os
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import MagicMock, patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from PySide6.QtCore import QDate, QEventLoop, QTimer, Qt
from PySide6.QtWidgets import QApplication, QDialog, QGroupBox, QHeaderView, QMessageBox

from core.auth import AuthenticationError
from core.db import Database
from core.services import PermitService
from gui.engedely_dialog import EngedelyDialog
from gui.workers import wait_for_tasks
import gui.main_window as main_ui
import gui.telephely_list_dialog as site_ui


class GuiTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.windows = []
        self.quit_patch = patch.object(self.app, "quit")
        self.quit = self.quit_patch.start()
        self.addCleanup(self.quit_patch.stop)

    def tearDown(self):
        for window in self.windows:
            window.close()
        wait_for_tasks()
        self.app.processEvents()
        for window in self.windows:
            window.deleteLater()

    def window(self, service, db=None, dirs=None):
        window = main_ui.MainWindow(db, dirs or {}, None, service)
        self.windows.append(window)
        return window

    def wait_until(self, predicate, timeout=3000):
        if predicate():
            return
        loop = QEventLoop()
        polling = QTimer()
        polling.setInterval(5)
        polling.timeout.connect(lambda: loop.quit() if predicate() else None)
        deadline = QTimer()
        deadline.setSingleShot(True)
        deadline.timeout.connect(loop.quit)
        polling.start()
        deadline.start(timeout)
        loop.exec()
        polling.stop()
        deadline.stop()
        self.assertTrue(predicate(), "Timed out waiting for the GUI worker")

    def accept_form(self, window, invoke, fill):
        failures = []

        def complete():
            dialog = None
            try:
                dialog = next(
                    item for item in window.findChildren(EngedelyDialog) if item.isVisible()
                )
                fill(dialog)
                dialog._on_accept()
                self.assertEqual(dialog.result(), QDialog.Accepted)
            except Exception as exc:
                failures.append(exc)
                if dialog is not None:
                    dialog.reject()

        with patch.object(main_ui, "error") as errors, patch.object(
            QMessageBox, "warning"
        ) as warnings:
            QTimer.singleShot(0, complete)
            invoke()
            self.assertFalse(failures, failures)
            self.assertFalse(errors.called, errors.call_args_list)
            self.assertFalse(warnings.called, warnings.call_args_list)


class MainWindowServiceTests(GuiTestCase):
    def setUp(self):
        super().setUp()
        self.workspace = tempfile.TemporaryDirectory(
            prefix=".main-window-test-", dir=Path(__file__).parent
        )
        self.root = Path(self.workspace.name)
        self.db = Database(self.root / "registry.db", self.root / "backups")
        self.service = PermitService(self.db, attachments_root=self.root / "attachments")
        self.service.bootstrap_admin("gui-admin", "gui-test-password", "GUI Admin")
        self.site = self.service.save_site({"nev": "A site"})
        self.other_site = self.service.save_site({"nev": "B site"})
        self.main = self.window(
            self.service, self.db, {"attachments": self.root / "attachments"}
        )

    def tearDown(self):
        try:
            super().tearDown()
        finally:
            self.service.logout()
            self.db.close()
            self.workspace.cleanup()

    def permit(self, **changes):
        values = {
            "telephely_id": self.site["id"],
            "engedely_nev": "Permit",
            "hatarozatszam": "TEST",
            "kategoria": main_ui.FIX_KATEGORIAK[0],
            "hataly_tipus": "fixed",
            "lejarat": "2090-01-10",
            "felulvizsgalatok": [
                {"type": "before_expiry", "days": 10, "note": "Keep note"}
            ],
        }
        values.update(changes)
        result = self.service.save_permit(values)
        self.assertTrue(self.main.refresh_data())
        return result

    def test_requires_an_authenticated_session(self):
        self.service.logout()
        with self.assertRaises(AuthenticationError):
            main_ui.MainWindow(self.db, {}, None, self.service)

    def test_snapshot_preserves_filters_selection_sort_and_last_good_state(self):
        first = self.permit()
        self.permit(engedely_nev="Second")
        self.main.telephely_filter.set_checked_by_data([self.site["id"]])
        self.main.table.sortItems(2, Qt.DescendingOrder)
        for row in range(self.main.table.rowCount()):
            if self.main.table.item(row, 0).data(Qt.UserRole)["id"] == first["id"]:
                self.main.table.selectRow(row)
        self.service.save_site(
            dict(self.site, nev="Renamed site"), expected_version=self.site["version"]
        )
        self.service.update_profile("Renamed User")
        self.assertTrue(self.main.refresh_data())
        self.assertEqual(self.main.telephely_filter.checked_data(), [self.site["id"]])
        self.assertEqual(
            [row["id"] for row in self.main._get_selected_entities()], [first["id"]]
        )
        self.assertEqual(
            self.main.table.horizontalHeader().sortIndicatorOrder(), Qt.DescendingOrder
        )
        self.assertTrue(
            all(
                row["responsible_name"] == "Renamed User"
                for row in self.main.data["engedelyek"]
            )
        )
        before = copy.deepcopy(self.main.data)
        with patch.object(
            self.service, "snapshot", side_effect=OSError("offline")
        ), patch.object(main_ui.logger, "error"):
            self.assertFalse(self.main.refresh_data())
        self.assertEqual(self.main.data, before)

    def test_actual_add_edit_forms_use_defaults_clean_payload_and_authorized_plan(self):
        self.main.telephely_filter.set_checked_by_data([self.other_site["id"]])
        self.main.kategoria_filter.set_checked_by_data([main_ui.FIX_KATEGORIAK[1]])
        source = self.root / "new.txt"
        source.write_text("New attachment")

        def create(dialog):
            self.assertEqual(dialog.telephely_combo.currentData(), self.other_site["id"])
            self.assertEqual(dialog.kategoria.currentText(), main_ui.FIX_KATEGORIAK[1])
            dialog.targy.setText("Created form")
            dialog.hatarozatszam.setText("CREATE")
            dialog.rb_fixed.setChecked(True)
            dialog.hatary_date.setDate(QDate(2090, 1, 10))
            dialog.felulvizsgalatok = [
                {"type": "before_expiry", "days": 10, "note": "Review note"}
            ]
            dialog.modosito_list.addItem("AMEND-1")
            dialog.attachment_plan.add(source)

        self.accept_form(self.main, self.main.add_engedely, create)
        self.assertEqual(len(self.main.data["engedelyek"]), 1)
        record = self.main.data["engedelyek"][0]
        self.assertEqual(record["modosito_hat"], "AMEND-1")
        self.assertEqual(record["_next_fv"].isoformat(), "2089-12-31")
        self.assertEqual(record["_fv_list"][0][1], "Review note")
        self.assertIn("created form", record["_search_index"])
        self.assertTrue((self.root / "attachments" / record["id"] / "new.txt").exists())

        def edit(dialog):
            dialog.targy.setText("Updated form")
            dialog.hatary_date.setDate(QDate(2090, 2, 10))
            dialog.modosito_list.addItem("AMEND-2")

        self.accept_form(self.main, lambda: self.main.edit_eng(record), edit)
        updated = self.main.data["engedelyek"][0]
        self.assertEqual(updated["version"], 2)
        self.assertEqual(updated["modosito_hat"], "AMEND-1;AMEND-2")
        self.assertEqual(updated["_next_fv"].isoformat(), "2090-01-31")
        self.assertIn("updated form", updated["_search_index"])
        self.main._user = dict(self.main._user, role="viewer")
        readonly = self.main._permit_dialog(updated)
        self.assertIsNotNone(readonly)
        self.assertFalse(readonly.save_btn.isEnabled())
        self.assertFalse(readonly.attach_add.isEnabled())
        readonly.reject()

    def test_bulk_versions_drop_and_soft_delete_preserve_files(self):
        record = self.permit()
        source = self.root / "attachment.txt"
        source.write_text("Preserve")
        with patch.object(main_ui, "error") as errors, patch.object(
            main_ui, "confirm", return_value=True
        ):
            self.main._add_files_to_entity(self.main.data["engedelyek"][0], [str(source)])
            self.assertEqual(self.main.data["engedelyek"][0]["version"], 2)
            self.main._bulk_archive(self.main.data["engedelyek"])
            self.assertEqual(self.main.data["engedelyek"][0]["archivalt"], 1)
            self.main._bulk_unarchive(self.main.data["engedelyek"])
            self.assertEqual(self.main.data["engedelyek"][0]["archivalt"], 0)
            self.main._bulk_delete(self.main.data["engedelyek"])
            self.assertEqual(self.main.data["engedelyek"], [])
            self.assertFalse(errors.called, errors.call_args_list)
        self.assertEqual(
            (self.root / "attachments" / record["id"] / "attachment.txt").read_text(),
            "Preserve",
        )

    def test_admin_assignment_uses_versions_and_unassigned_filter(self):
        self.permit()
        self.main.table.selectRow(0)
        with patch.object(
            main_ui.QInputDialog, "getItem", return_value=("Nincs felelős", True)
        ), patch.object(main_ui, "error") as errors:
            self.main.assign_selected()
            self.assertFalse(errors.called, errors.call_args_list)
        record = self.main.data["engedelyek"][0]
        self.assertIsNone(record["responsible_user_id"])
        self.assertEqual(record["version"], 2)
        self.assertEqual(record["responsible_name"], "Nincs felelős")
        self.main.responsibility_filter.setCurrentIndex(
            self.main.responsibility_filter.findData("unassigned")
        )
        self.assertEqual(len(self.main._get_filtered_items()), 1)

    def test_site_dialog_crud_versions_and_failed_refresh_preservation(self):
        dialog = site_ui.TelephelyListDialog(self.service, self.main)
        editor = MagicMock()
        editor.exec.return_value = QDialog.Accepted
        editor.get_data.return_value = {"nev": "New site", "ktj": None, "cim": None}
        with patch.object(site_ui, "TelephelyDialog", return_value=editor), patch.object(
            QMessageBox, "question", return_value=QMessageBox.Yes
        ):
            dialog.new_telephely()
            self.assertEqual(dialog.listw.count(), 3)
            for row in range(dialog.listw.count()):
                if dialog.listw.item(row).data(Qt.UserRole)["nev"] == "New site":
                    dialog.listw.setCurrentRow(row)
            editor.get_data.return_value = {
                "nev": "Updated site", "ktj": "123", "cim": "Address"
            }
            dialog.edit_selected()
            self.assertEqual(dialog.listw.currentItem().data(Qt.UserRole)["version"], 2)
            with patch.object(
                self.service, "snapshot", side_effect=OSError("offline")
            ), patch.object(QMessageBox, "critical"):
                dialog.refresh_list()
                self.assertEqual(dialog.listw.count(), 3)
            dialog.delete_selected()
            self.assertEqual(dialog.listw.count(), 2)
        dialog.close()

    def test_password_change_logout_order_and_normal_close(self):
        self.main.show()
        child = QDialog(self.main)
        child.show()
        events = []
        self.main.logout_requested.connect(
            lambda: events.append((
                self.main.isVisible(), child.isVisible(),
                self.main.timer.isActive(), self.service.session,
            ))
        )
        dialog = MagicMock()
        dialog.exec.side_effect = lambda: self.service.change_password(
            "gui-test-password", "new-test-password", display_name="New Display"
        )
        with patch("gui.auth_dialogs.ProfileDialog", return_value=dialog):
            self.main.open_profile()
        self.assertEqual(events, [(False, False, False, None)])
        self.assertFalse(self.quit.called)
        self.service.login("gui-admin", "new-test-password")
        self.assertEqual(self.service.current_user["display_name"], "New Display")
        normal = self.window(self.service, self.db)
        normal.close()
        self.assertEqual(self.quit.call_count, 1)

    def test_backup_worker_keeps_gui_responsive(self):
        self.main.show()
        responsive = []

        def slow_backup():
            time.sleep(0.12)
            return self.root / "backup.db"

        with patch.object(
            self.service, "create_backup", side_effect=slow_backup
        ), patch.object(main_ui, "info") as info:
            self.main.create_backup_now()
            QTimer.singleShot(20, lambda: responsive.append(self.main._backup_running))
            self.wait_until(lambda: not self.main._backup_running)
            self.assertEqual(responsive, [True])
            self.assertEqual(info.call_count, 1)


class FakeAuth:
    def __init__(self):
        self.session = object()
        self.current_user = {
            "id": "own", "username": "admin", "display_name": "Owner",
            "role": "admin", "active": 1,
        }

    def logout(self):
        self.session = None


class FakeService:
    def __init__(self):
        self.auth = FakeAuth()
        self.users = [
            dict(self.auth.current_user),
            {"id": "other", "username": "clerk", "display_name": "Other Person",
             "role": "clerk", "active": 1},
            {"id": "historic", "username": "retired", "display_name": "Historical Person",
             "role": "clerk", "active": 0},
        ]
        self.records = [
            {"id": "p-" + name, "telephely_id": "site",
             "responsible_user_id": assignee, "version": 1}
            for name, assignee in (
                ("own", "own"), ("other", "other"),
                ("historic", "historic"), ("unassigned", None),
            )
        ]
        self.count_delay = 0
        self.count_calls = []
        self.count_threads = []
        self.count_active = 0
        self.max_count_active = 0
        self.snapshot_delay = 0
        self.snapshot_failure = None
        self.snapshot_started = threading.Event()

    def snapshot(self):
        snapshot = {
            "current_user": dict(self.auth.current_user),
            "users": copy.deepcopy(self.users),
            "telephelyek": [{"id": "site", "nev": "Site"}],
            "engedelyek": copy.deepcopy(self.records),
        }
        delay, failure = self.snapshot_delay, self.snapshot_failure
        self.snapshot_started.set()
        time.sleep(delay)
        if failure is not None:
            raise failure
        return snapshot

    def count_attachments(self, identifiers):
        self.count_calls.append(tuple(identifiers))
        self.count_threads.append(threading.get_ident())
        self.count_active += 1
        self.max_count_active = max(self.max_count_active, self.count_active)
        try:
            time.sleep(self.count_delay)
            return {identifier: 2 for identifier in identifiers}
        finally:
            self.count_active -= 1


class MainWindowUiTests(GuiTestCase):
    def setUp(self):
        super().setUp()
        self.service = FakeService()
        self.main = self.window(self.service)
        self.wait_until(lambda: not self.main._attachment_counts_running)

    def select_person(self, value):
        index = self.main.responsibility_filter.findData(value)
        self.assertGreaterEqual(index, 0)
        self.main.responsibility_filter.setCurrentIndex(index)

    def filtered_ids(self):
        return [row["id"] for row in self.main._get_filtered_items()]

    def export_directory(self):
        workspace = tempfile.TemporaryDirectory(
            prefix=".main-export-test-", dir=Path(__file__).parent
        )
        self.addCleanup(workspace.cleanup)
        return Path(workspace.name)

    def test_arbitrary_person_filter_and_one_click_shortcuts(self):
        for value, expected in (
            ("person:other", "p-other"), ("own", "p-own"),
            ("unassigned", "p-unassigned"),
        ):
            self.select_person(value)
            self.assertEqual(self.filtered_ids(), [expected])
            self.assertEqual(self.main.table.rowCount(), 1)

    def test_import_does_not_replace_application_stylesheet(self):
        result = subprocess.run(
            [
                sys.executable, "-c",
                "from PySide6.QtWidgets import QApplication; "
                "app=QApplication([]); "
                "app.setStyleSheet('QWidget { color: red; }'); "
                "import gui.main_window; "
                "assert app.styleSheet() == 'QWidget { color: red; }'",
            ],
            cwd=Path(__file__).resolve().parents[1],
            env=dict(os.environ, QT_QPA_PLATFORM="offscreen"),
            capture_output=True,
            text=True,
            timeout=15,
        )
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_disabled_historical_person_remains_named_and_filterable(self):
        self.select_person("person:historic")
        self.assertIn("Historical Person", self.main.responsibility_filter.currentText())
        self.assertIn("inaktív", self.main.responsibility_filter.currentText())
        self.assertEqual(self.filtered_ids(), ["p-historic"])
        self.assertEqual(self.main.table.item(0, 9).text(), "Historical Person")

    def test_person_selection_survives_snapshot_rename_and_missing_historical_user(self):
        self.select_person("person:historic")
        self.service.users[-1]["display_name"] = "Historical Renamed"
        self.main.refresh_data()
        self.assertEqual(self.main.responsibility_filter.currentData(), "person:historic")
        self.assertIn("Historical Renamed", self.main.responsibility_filter.currentText())
        self.service.users = [u for u in self.service.users if u["id"] != "historic"]
        self.main.refresh_data()
        self.assertEqual(self.main.responsibility_filter.currentData(), "person:historic")
        self.assertEqual(self.filtered_ids(), ["p-historic"])

    def test_revoked_session_during_profile_construction_logs_out(self):
        with patch(
            "gui.auth_dialogs.ProfileDialog", side_effect=AuthenticationError("revoked")
        ):
            self.main.open_profile()
        self.assertTrue(self.main._logged_out)
        self.assertIsNone(self.service.auth.session)
        self.assertEqual(self.main.responsibility_filter.count(), 0)
        self.assertFalse(self.main.timer.isActive())

    def test_revoked_session_during_profile_exec_logs_out(self):
        dialog = MagicMock()
        dialog.exec.side_effect = AuthenticationError("revoked")
        with patch("gui.auth_dialogs.ProfileDialog", return_value=dialog):
            self.main.open_profile()
        self.assertTrue(self.main._logged_out)
        self.assertIsNone(self.service.auth.session)

    def test_committed_mutation_surfaces_cleanup_warnings_as_success(self):
        committed = []
        message = "Az adatbázis mentése sikerült; helyreállítás: recovery-directory"
        with patch.object(
            self.service, "take_warnings", return_value=[message], create=True
        ), patch.object(self.main, "refresh_data", return_value=True) as refresh, patch.object(
            main_ui, "warning"
        ) as warning, patch.object(main_ui, "error") as error:
            self.assertTrue(self.main._mutate(lambda: committed.append(True)))
        self.assertEqual(committed, [True])
        self.assertEqual(refresh.call_count, 1)
        self.assertFalse(error.called)
        warning.assert_called_once_with(
            self.main, "Művelet végrehajtva – figyelmeztetés", message
        )

    def test_warning_lookup_failure_does_not_report_committed_mutation_as_failed(self):
        committed = []
        with patch.object(
            self.service, "take_warnings", side_effect=OSError("offline"), create=True
        ), patch.object(self.main, "refresh_data", return_value=True), patch.object(
            main_ui, "warning"
        ) as warning, patch.object(main_ui, "error") as error:
            self.assertTrue(self.main._mutate(lambda: committed.append(True)))
        self.assertEqual(committed, [True])
        self.assertFalse(error.called)
        self.assertIn("mentése sikerült", warning.call_args.args[2])

    def test_revocation_after_committed_mutation_logs_out_without_false_failure(self):
        committed = []
        with patch.object(
            self.service, "take_warnings",
            side_effect=AuthenticationError("revoked"), create=True,
        ), patch.object(main_ui, "error") as error:
            self.assertTrue(self.main._mutate(lambda: committed.append(True)))
        self.assertEqual(committed, [True])
        self.assertTrue(self.main._logged_out)
        self.assertIsNone(self.service.auth.session)
        self.assertFalse(error.called)

    def test_attachment_counts_are_async_and_rendering_does_not_reschedule(self):
        self.service.count_delay = 0.12
        self.main.refresh_data()
        self.assertEqual(self.main.table.item(0, 8).text(), "Csatolmányok…")
        responsive = []
        QTimer.singleShot(
            20, lambda: responsive.append(self.main._attachment_counts_running)
        )
        self.wait_until(lambda: not self.main._attachment_counts_running)
        self.assertEqual(responsive, [True])
        self.assertEqual(self.main.table.item(0, 8).text(), "Megnyitás (2)")
        self.assertNotEqual(self.service.count_threads[-1], threading.get_ident())
        before = len(self.service.count_calls)
        for _ in range(5):
            self.main.refresh_table()
        self.assertEqual(len(self.service.count_calls), before)

    def test_stale_attachment_count_generation_is_discarded(self):
        self.service.count_delay = 0.08
        self.main.refresh_data()
        self.service.records = [{"id": "latest", "telephely_id": "site", "version": 1}]
        self.main.refresh_data()
        self.wait_until(
            lambda: not self.main._attachment_counts_running
            and self.main._pending_attachment_counts is None
        )
        self.assertEqual(self.main._attachments_cache, {"latest": 2})
        self.assertEqual(self.service.max_count_active, 1)

    def test_logout_discards_attachment_count_results(self):
        self.service.count_delay = 0.08
        self.main.refresh_data()
        self.main.logout()
        self.wait_until(lambda: not self.main._attachment_counts_running)
        self.assertEqual(self.main._attachments_cache, {})
        self.assertEqual(self.main.table.rowCount(), 0)

    def test_logout_hides_busy_dialog_before_emitting_signal(self):
        self.main.show()
        child = QDialog(self.main)
        child.show()
        events = []
        self.main.logout_requested.connect(
            lambda: events.append((
                self.main.isVisible(), child.isVisible(), self.service.auth.session,
            ))
        )
        with patch.object(child, "reject"), patch.object(child, "close", return_value=False):
            self.main.logout()
        self.assertEqual(events, [(False, False, None)])

    def test_async_snapshot_preserves_view_and_applies_latest_queued_generation(self):
        self.service.snapshot_delay = 0.08
        self.service.snapshot_started.clear()
        self.service.records = [{"id": "intermediate", "telephely_id": "site", "version": 1}]
        self.main.manual_refresh()
        self.assertTrue(self.service.snapshot_started.wait(1))
        self.assertEqual(len(self.main.data["engedelyek"]), 4)
        self.service.records = [{"id": "latest", "telephely_id": "site", "version": 1}]
        self.main.manual_refresh()
        self.wait_until(
            lambda: not self.main._snapshot_refresh_running
            and self.main._pending_snapshot_refresh is None
        )
        self.assertEqual(self.main.data["engedelyek"][0]["id"], "latest")

    def test_async_snapshot_failure_preserves_last_good_state(self):
        before = copy.deepcopy(self.main.data)
        self.service.snapshot_failure = OSError("offline")
        with patch.object(main_ui.logger, "error"):
            self.main.manual_refresh()
            self.wait_until(lambda: not self.main._snapshot_refresh_running)
        self.assertEqual(self.main.data, before)
        self.assertIn("utolsó sikeres", self.main.statusBar().currentMessage())

    def test_logout_discards_running_snapshot(self):
        self.service.snapshot_delay = 0.08
        self.main.manual_refresh()
        self.main.logout()
        self.wait_until(lambda: not self.main._snapshot_refresh_running)
        self.assertEqual(self.main.data["engedelyek"], [])
        self.assertEqual(self.main.table.rowCount(), 0)

    def test_familiar_grouped_controls_and_table_dimensions_are_preserved(self):
        groups = self.main.findChildren(QGroupBox)
        self.assertEqual(
            [group.title() for group in groups],
            ["Adminisztráció", "Exportálás", "Rendszer / Mentések"],
        )
        bottom = self.main.centralWidget().layout().itemAt(
            self.main.centralWidget().layout().count() - 1
        ).layout()
        self.assertEqual(
            [bottom.itemAt(i).widget() for i in range(bottom.count())
             if isinstance(bottom.itemAt(i).widget(), QGroupBox)],
            groups,
        )
        self.assertEqual(self.main.btn_assign.parentWidget().title(), "Adminisztráció")
        self.assertEqual(self.main.btn_refresh.parentWidget().title(), "Rendszer / Mentések")
        header = self.main.table.horizontalHeader()
        self.assertEqual(header.sectionResizeMode(7), QHeaderView.Stretch)
        self.assertEqual(
            [self.main.table.columnWidth(c) for c in (2, 5, 6, 8)], [320, 78, 97, 105]
        )

    def test_batch_pdf_keeps_dialogs_on_gui_and_exports_copied_data_off_thread(self):
        target = self.export_directory() / "permits.pdf"
        self.select_person("own")
        record = self.main.data["engedelyek"][0]
        record.update(engedely_nev="Original", hatarozatszam="PDF",
                      kategoria=main_ui.FIX_KATEGORIAK[0])
        original_exporter = main_ui.export_groups_pdf
        started = threading.Event()
        threads, exported, responsive = [], [], []

        def choose(*_):
            threads.append(("selection", threading.get_ident()))
            return ["Site"]

        def choose_path(*_):
            threads.append(("save", threading.get_ident()))
            return str(target), ""

        def export(path, groups, names, **kwargs):
            threads.append(("export", threading.get_ident()))
            exported.append(copy.deepcopy(groups))
            started.set()
            time.sleep(0.12)
            return original_exporter(path, groups, names, **kwargs)

        with patch.object(self.main, "select_items_popup", side_effect=choose), patch.object(
            main_ui.QFileDialog, "getSaveFileName", side_effect=choose_path
        ), patch.object(main_ui, "export_groups_pdf", side_effect=export), patch.object(
            main_ui, "info", side_effect=lambda *_: threads.append(("success", threading.get_ident()))
        ), patch.object(main_ui, "error") as errors:
            self.main.batch_export("telephely")
            self.assertTrue(started.wait(1))
            record["engedely_nev"] = "Changed after dispatch"
            QTimer.singleShot(20, lambda: responsive.append(self.main._export_running))
            self.wait_until(lambda: not self.main._export_running)
            self.assertFalse(errors.called, errors.call_args_list)
        self.assertEqual(responsive, [True])
        self.assertEqual(exported[0]["Site"][0]["engedely_nev"], "Original")
        self.assertTrue(target.read_bytes().startswith(b"%PDF"))
        for kind, thread in threads:
            if kind == "export":
                self.assertNotEqual(thread, threading.get_ident())
            else:
                self.assertEqual(thread, threading.get_ident())

    def test_excel_export_is_async_and_preserves_filters_public_names_and_archived_option(self):
        target = self.export_directory() / "permits.xlsx"
        self.select_person("own")
        self.main.show_archived.setChecked(True)
        record = self.main.data["engedelyek"][0]
        record.update(engedely_nev="Original", hatarozatszam="XLSX", archivalt=1,
                      kategoria=main_ui.FIX_KATEGORIAK[0])
        original_exporter = main_ui.export_excel_fn
        captured = []

        def export(path, records, sites, **kwargs):
            captured.append((threading.get_ident(), copy.deepcopy(records), kwargs))
            time.sleep(0.08)
            return original_exporter(path, records, sites, **kwargs)

        with patch.object(
            main_ui.QFileDialog, "getSaveFileName", return_value=(str(target), "")
        ), patch.object(main_ui, "export_excel_fn", side_effect=export), patch.object(
            main_ui, "info"
        ), patch.object(main_ui, "error") as errors:
            self.main.export_to_excel()
            self.assertTrue(self.main._export_running)
            self.assertFalse(self.main.btn_excel.isEnabled())
            record["engedely_nev"] = "Changed after dispatch"
            self.wait_until(lambda: not self.main._export_running)
            self.assertFalse(errors.called, errors.call_args_list)
        self.assertTrue(self.main.btn_excel.isEnabled())
        self.assertNotEqual(captured[0][0], threading.get_ident())
        self.assertEqual(len(captured[0][1]), 1)
        self.assertEqual(captured[0][1][0]["engedely_nev"], "Original")
        self.assertTrue(captured[0][2]["options"]["include_archived"])
        workbook = load_workbook(target)
        try:
            values = [value for row in workbook.active.iter_rows(values_only=True) for value in row]
            self.assertIn("Owner", values)
            self.assertIn("Original", values)
        finally:
            workbook.close()

    def test_export_authentication_failure_logs_out_before_writing(self):
        target = self.export_directory() / "denied.xlsx"
        with patch.object(
            main_ui.QFileDialog, "getSaveFileName", return_value=(str(target), "")
        ), patch.object(self.service, "snapshot", side_effect=AuthenticationError("revoked")), patch.object(
            main_ui, "export_excel_fn"
        ) as exporter, patch.object(main_ui, "info") as info, patch.object(
            main_ui, "error"
        ) as error:
            self.main.export_to_excel()
            self.wait_until(lambda: not self.main._export_running)
        self.assertTrue(self.main._logged_out)
        self.assertFalse(exporter.called)
        self.assertFalse(info.called)
        self.assertFalse(error.called)
        self.assertFalse(target.exists())

    def test_export_completion_after_logout_does_not_open_message_boxes(self):
        target = self.export_directory() / "ignored.xlsx"
        started = threading.Event()

        def export(*_, **__):
            started.set()
            time.sleep(0.08)

        with patch.object(
            main_ui.QFileDialog, "getSaveFileName", return_value=(str(target), "")
        ), patch.object(main_ui, "export_excel_fn", side_effect=export), patch.object(
            main_ui, "info"
        ) as info, patch.object(main_ui, "error") as error:
            self.main.export_to_excel()
            self.assertTrue(started.wait(1))
            self.main.logout()
            self.wait_until(lambda: not self.main._export_running)
        self.assertFalse(info.called)
        self.assertFalse(error.called)

    def test_export_failure_restores_controls_and_reports_error_on_gui_thread(self):
        target = self.export_directory() / "failed.xlsx"
        threads = []
        with patch.object(
            main_ui.QFileDialog, "getSaveFileName", return_value=(str(target), "")
        ), patch.object(main_ui, "export_excel_fn", side_effect=OSError("disk full")), patch.object(
            main_ui, "info"
        ) as info, patch.object(
            main_ui, "error", side_effect=lambda *_: threads.append(threading.get_ident())
        ) as error:
            self.main.export_to_excel()
            self.wait_until(lambda: not self.main._export_running)
        self.assertFalse(info.called)
        self.assertEqual(error.call_count, 1)
        self.assertEqual(threads, [threading.get_ident()])
        self.assertTrue(self.main.btn_excel.isEnabled())
        self.assertTrue(self.main.btn_quick_tp.isEnabled())
        self.assertTrue(self.main.btn_quick_kat.isEnabled())

    def test_legacy_executable_and_hidden_attachment_open_requires_confirmation(self):
        record = self.main.data["engedelyek"][0]
        root = self.export_directory()
        for filename in ("legacy.exe", "legacy.CMD", ".hidden", "legacy.desktop"):
            with self.subTest(filename=filename), patch.object(
                self.service, "list_attachments", return_value=[filename], create=True
            ), patch.object(
                self.service, "attachment_path", return_value=root / filename, create=True
            ), patch.object(main_ui, "confirm", return_value=False) as confirm, patch.object(
                main_ui, "_open_path"
            ) as open_path:
                self.main._open_with_preference(record)
                self.assertEqual(confirm.call_count, 1)
                self.assertFalse(open_path.called)

    def test_confirmed_legacy_attachment_is_revalidated_before_opening(self):
        record = self.main.data["engedelyek"][0]
        path = self.export_directory() / "legacy.exe"
        with patch.object(
            self.service, "list_attachments", return_value=[path.name], create=True
        ), patch.object(
            self.service, "attachment_path", return_value=path, create=True
        ) as attachment_path, patch.object(
            main_ui, "confirm", return_value=True
        ), patch.object(main_ui, "_open_path") as open_path:
            self.main._open_with_preference(record)
        self.assertEqual(attachment_path.call_count, 2)
        open_path.assert_called_once_with(path)

    def test_supported_attachment_opens_without_legacy_warning(self):
        record = self.main.data["engedelyek"][0]
        path = self.export_directory() / "document.pdf"
        with patch.object(
            self.service, "list_attachments", return_value=[path.name], create=True
        ), patch.object(
            self.service, "attachment_path", return_value=path, create=True
        ), patch.object(main_ui, "confirm") as confirm, patch.object(
            main_ui, "_open_path"
        ) as open_path:
            self.main._open_with_preference(record)
        self.assertFalse(confirm.called)
        open_path.assert_called_once_with(path)

    def test_multiple_attachments_open_folder_without_executing_a_file(self):
        record = self.main.data["engedelyek"][0]
        root = self.export_directory()
        with patch.object(
            self.service, "list_attachments",
            return_value=["legacy.exe", "document.pdf"], create=True,
        ), patch.object(
            self.service, "attachment_path", return_value=root / "legacy.exe", create=True
        ), patch.object(main_ui, "confirm") as confirm, patch.object(
            main_ui, "_open_path"
        ) as open_path:
            self.main._open_with_preference(record)
        self.assertFalse(confirm.called)
        open_path.assert_called_once_with(root)

    def test_revocation_during_legacy_warning_prevents_open_and_logs_out(self):
        record = self.main.data["engedelyek"][0]
        path = self.export_directory() / "legacy.exe"
        with patch.object(
            self.service, "list_attachments", return_value=[path.name], create=True
        ), patch.object(
            self.service, "attachment_path",
            side_effect=[path, AuthenticationError("revoked")], create=True,
        ), patch.object(main_ui, "confirm", return_value=True), patch.object(
            main_ui, "_open_path"
        ) as open_path, patch.object(main_ui, "error") as error:
            self.main._open_with_preference(record)
        self.assertFalse(open_path.called)
        self.assertFalse(error.called)
        self.assertTrue(self.main._logged_out)
        self.assertIsNone(self.service.auth.session)


class ReviewExportSafetyTests(GuiTestCase):
    def setUp(self):
        super().setUp()
        self.workspace = tempfile.TemporaryDirectory(
            prefix=".review-export-test-", dir=Path(__file__).parent
        )
        self.root = Path(self.workspace.name)
        self.notes = [
            "=2+3", "+2+3", "-1+2", "@SUM(A1)",
            " \t=2+3", "\r\n+2", "\t@SUM(A1)", "ordinary note",
        ]
        record = {
            "telephely_id": "site", "engedely_nev": "+PERMIT()",
            "hatarozatszam": "TEST", "kategoria": "Category",
        }
        self.dialog = main_ui.ReviewDialog(
            [(record, date(2030, 1, i + 1), note) for i, note in enumerate(self.notes)],
            [{"id": "site", "nev": "=SITE()"}],
        )

    def tearDown(self):
        self.dialog.close()
        self.dialog.deleteLater()
        super().tearDown()
        self.workspace.cleanup()

    def export(self, filename):
        target = self.root / filename
        with patch.object(
            main_ui.QFileDialog, "getSaveFileName", return_value=(str(target), "")
        ), patch.object(main_ui, "info"), patch.object(main_ui, "error") as errors:
            self.dialog._export_excel()
            self.assertFalse(errors.called, errors.call_args_list)
        return target

    def test_review_csv_neutralizes_formulas_in_names_notes_and_whitespace(self):
        path = self.export("reviews.csv")
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream, delimiter=";"))
        for row, note in zip(rows[1:], self.notes):
            self.assertEqual(row[1], "'=SITE()")
            self.assertEqual(row[2], "'+PERMIT()")
            expected = "'" + note if note.lstrip().startswith(("=", "+", "-", "@")) else note
            self.assertEqual(row[3], expected)

    def test_review_xlsx_text_cells_are_not_formulas(self):
        workbook = load_workbook(self.export("reviews.xlsx"), data_only=False)
        try:
            sheet = workbook.active
            for row in sheet.iter_rows(min_row=2):
                self.assertTrue(all(cell.data_type == "s" for cell in row))
                self.assertEqual(row[1].value, "=SITE()")
                self.assertEqual(row[2].value, "+PERMIT()")
            self.assertEqual(sheet.cell(2, 4).value, "=2+3")
        finally:
            workbook.close()

    def test_safe_csv_text_is_unchanged(self):
        for text in ("2030-01-01", "ordinary text", "  ordinary text", "", "Permit-123"):
            self.assertEqual(self.dialog._csv_text(text), text)


if __name__ == "__main__":
    unittest.main()
