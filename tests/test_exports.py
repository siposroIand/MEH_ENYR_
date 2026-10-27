import importlib
import json
from datetime import date
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from exporters import excel_exporter, pdf_exporter, report_builder_exporter
from utils.dates import compute_review_date


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.item = {
            "id": "permit",
            "telephely_id": "site",
            "kategoria": "Hulladékgazdálkodási engedély",
            "engedely_nev": "Permit <subject> & details",
            "kiado_hatosag": "Authority",
            "hatarozatszam": "A/123",
            "modosito_hat": "B/1; C/2",
            "hataly_tipus": "fixed",
            "lejarat": "2027-06-30",
            "letrejott": "2026-01-01",
            "responsible_name": "Responsible Example",
            "responsible_user_id": "private-user-id",
            "password_hash": "private-password-value",
            "email": "private-email-value",
            "users": {"password": "private-nested-value"},
            "megjegyzes": "Keep <literal> & notes",
            "hulladek_nv_elokez": "15",
            "felulvizsgalat_date": "2020-01-01",
            "felulvizsgalat_nap": 99,
            "felulvizsgalatok": [
                {"type": "before_expiry", "days": 0, "note": "Zero offset"},
                {"date": "2027-01-02", "note": "Legacy note"},
                {"type": "before_expiry", "days": 30, "note": "Typed <note> & data"},
                {"type": "date", "date": "2027-01-02", "note": "Second same-date note"},
            ],
        }
        self.sites = [{"id": "site", "nev": "Example Site"}]

    def test_imports_headlessly(self):
        for module in (excel_exporter, pdf_exporter, report_builder_exporter):
            self.assertIs(importlib.import_module(module.__name__), module)

    def test_shared_parser_retains_all_dates_and_notes(self):
        expected = [
            (date(2027, 1, 2), "Legacy note"),
            (date(2027, 1, 2), "Second same-date note"),
            (date(2027, 5, 31), "Typed <note> & data"),
            (date(2027, 6, 30), "Zero offset"),
        ]
        for module in (excel_exporter, pdf_exporter, report_builder_exporter):
            for raw in (self.item["felulvizsgalatok"], json.dumps(self.item["felulvizsgalatok"])):
                with self.subTest(module=module.__name__, raw_type=type(raw).__name__):
                    self.assertEqual(module._extract_review_entries(
                        dict(self.item, felulvizsgalatok=raw)), expected)
        self.assertEqual(compute_review_date(self.item), date(2027, 1, 2))

    def test_explicit_empty_reviews_clear_stale_legacy_values(self):
        for raw in ([], "[]"):
            item = dict(self.item, felulvizsgalatok=raw)
            self.assertIsNone(compute_review_date(item))
            for module in (excel_exporter, pdf_exporter, report_builder_exporter):
                self.assertEqual(module._extract_review_entries(item), [])

    def test_legacy_and_invalid_reviews(self):
        item = dict(self.item, felulvizsgalatok=None)
        self.assertEqual(compute_review_date(item), date(2020, 1, 1))
        item["felulvizsgalat_date"] = None
        self.assertEqual(compute_review_date(item), date(2027, 3, 23))
        for raw in ("invalid json", "{}", [{"type": "unknown"}]):
            self.assertEqual(excel_exporter._extract_review_entries(
                dict(self.item, felulvizsgalatok=raw)), [])

    def test_legacy_dotted_and_datetime_review_formats(self):
        for value in ("2027.01.02", "2027-01-02T10:30:00"):
            item = dict(self.item, felulvizsgalatok=[{"date": value, "note": "Original note"}])
            for module in (excel_exporter, pdf_exporter, report_builder_exporter):
                self.assertEqual(module._extract_review_entries(item),
                                 [(date(2027, 1, 2), "Original note")])
            self.assertEqual(compute_review_date(dict(
                self.item, felulvizsgalatok=None, felulvizsgalat_date=value,
                felulvizsgalat_nap=None)), date(2027, 1, 2))

    def test_real_excel_output(self):
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            path = Path(directory) / "permits.xlsx"
            excel_exporter.export_to_excel(path, [self.item], self.sites)
            workbook = load_workbook(path, rich_text=True)
            try:
                sheet = workbook["Összes"]
                headers = [cell.value for cell in sheet[1]]
                row = {header: str(sheet.cell(2, index + 1).value)
                       for index, header in enumerate(headers)}
                self.assertEqual(row["Felelős"], "Responsible Example")
                self.assertEqual(row["Megjegyzés"], self.item["megjegyzes"])
                self.assertIn("2027.06.30: Zero offset", row["Felülvizsgálat"])
                self.assertIn("Second same-date note", row["Felülvizsgálat"])
                rich_review = sheet.cell(2, headers.index("Felülvizsgálat") + 1).value
                self.assertTrue(any(getattr(getattr(part, "font", None), "b", False)
                                    for part in rich_review))
                self.assertNotIn("private-", str(row))
            finally:
                workbook.close()

    def test_excel_filtered_empty_is_clear_error(self):
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            with self.assertRaisesRegex(ValueError, "No items"):
                excel_exporter.export_to_excel(
                    Path(directory) / "empty.xlsx", [dict(self.item, archivalt=1)], self.sites)

    def test_pdf_output_and_retained_content(self):
        captured = []
        original_table = pdf_exporter.Table

        def table(data, *args, **kwargs):
            captured.extend(cell.getPlainText() for row in data for cell in row
                            if hasattr(cell, "getPlainText"))
            return original_table(data, *args, **kwargs)

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            path = Path(directory) / "permits.pdf"
            with patch.object(pdf_exporter, "Table", side_effect=table):
                pdf_exporter.export_groups_pdf(
                    str(path), {"Example Site": [self.item]}, {"site": "Example Site"},
                    options={"include_logo": False})
            self.assertTrue(path.read_bytes().startswith(b"%PDF-"))
            self.assertGreater(path.stat().st_size, 1000)
        text = "\n".join(captured)
        for value in ("Responsible Example", "Legacy note", "Second same-date note",
                      "Typed <note> & data", "Zero offset", "NV-Elő: 15"):
            self.assertIn(value, text)
        self.assertNotIn("private-", text)

    def test_real_reports_and_private_column_allowlist(self):
        columns = ["telephely", "engedely_nev", "responsible_name", "felulvizsgalat",
                   "megjegyzes", "responsible_user_id", "password_hash", "email", "users"]
        row = report_builder_exporter._prepare_row(self.item, {"site": "Example Site"}, columns)
        self.assertNotIn("private-", str(row))
        self.assertIn("Zero offset", row["felulvizsgalat"])
        self.assertIn("Legacy note", row["megjegyzes"])
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            for fmt, suffix in (("PDF", "pdf"), ("XLSX", "xlsx")):
                path = Path(directory) / f"report.{suffix}"
                report_builder_exporter.export_report(
                    path, [self.item], self.sites, columns, fmt=fmt, title="Example report")
                self.assertGreater(path.stat().st_size, 1000)
                if fmt == "PDF":
                    self.assertTrue(path.read_bytes().startswith(b"%PDF-"))
                else:
                    workbook = load_workbook(path)
                    try:
                        values = list(workbook.active.values)
                        self.assertIn("Felelős", values[0])
                        self.assertNotIn("responsible_user_id", values[0])
                        self.assertIn("Responsible Example", values[1])
                        self.assertNotIn("private-", str(values))
                    finally:
                        workbook.close()

    def test_report_per_group_wrapper(self):
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            for fmt, suffix in (("PDF", "pdf"), ("XLSX", "xlsx")):
                report_builder_exporter.export_report(
                    directory, {"Site One": [self.item], "Site Two": [self.item]},
                    self.sites, ["engedely_nev", "responsible_name", "felulvizsgalat"],
                    fmt=fmt, per_group=True)
                self.assertTrue((Path(directory) / f"Site_One.{suffix}").is_file())
                self.assertTrue((Path(directory) / f"Site_Two.{suffix}").is_file())


if __name__ == "__main__":
    unittest.main()
